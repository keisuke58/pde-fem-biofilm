"""make_slides_3d_demo_pptx.py -- PowerPoint version of slides_3d_demo.tex with the videos embedded (they play inside
the slide; the still of the last frame is the poster). Same content as the beamer deck.

    python make_slides_3d_demo_pptx.py   -> slides_3d_demo.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
A = ROOT / "assets"
BLUE = RGBColor(30, 80, 160)
FONT = "Times New Roman"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=18, color=None, bold=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, ln in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln
        p.space_after = Pt(6)
        for r in p.runs:
            r.font.size = Pt(size); r.font.name = FONT; r.font.bold = bold
            if color is not None:
                r.font.color.rgb = color
    return tb


def frame(title, n):
    s = prs.slides.add_slide(blank)
    bar = s.shapes.add_shape(1, 0, 0, W, Inches(0.85))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    text(s, Inches(0.4), Inches(0.12), Inches(12.5), Inches(0.7), title, size=26, color=RGBColor(255, 255, 255))
    text(s, Inches(11.8), Inches(7.05), Inches(1.4), Inches(0.4), f"{n}", size=12, color=RGBColor(120, 120, 120))
    return s


def movie(s, name, x, y, w):
    png = A / f"{name}.png"
    from PIL import Image
    iw, ih = Image.open(png).size
    h = Emu(int(w * ih / iw))
    s.shapes.add_movie(str(A / f"{name}.mp4"), x, y, w, h, poster_frame_image=str(png), mime_type="video/mp4")
    return h


# 1 title
s = prs.slides.add_slide(blank)
text(s, Inches(1), Inches(2.3), Inches(11.3), Inches(1), "Biofilm Growth and Stress on a Real Molar", size=40, color=BLUE,
     bold=True)
text(s, Inches(1), Inches(3.3), Inches(11.3), Inches(1),
     "3D demonstrations of the growth model, with converged stresses and no experimental data", size=20)
text(s, Inches(1), Inches(4.6), Inches(11.3), Inches(1.5),
     ["Keisuke Nishioka", "Institute of Continuum Mechanics (IKM), Leibniz Universität Hannover", "20 October 2026"],
     size=18)

s = frame("What I show", 2)
text(s, Inches(0.6), Inches(1.2), Inches(12.1), Inches(5.8), [
    "• Growth model of Klempt et al. (2024): phase φ with dφ/dt = β∇²φ + k_α α (plus the bound on φ), growth F_g = α I, "
    "stiffness E(φ² + f), E = 10 Pa, ν = 0.49.",
    "• Added in this work: a nutrient supplied by the saliva and used up by the biofilm (Klempt et al. 2024, Eq. 35), "
    "and a growth limit by pressure, ω = min(1, max(0, 1 − p/p_h)).",
    "• Geometry: the thesis's 2 mm cube for verification, then a real lower first molar at its real size "
    "(surface from a public CC BY 4.0 data set).",
    "• Every stress I quote is checked against mesh refinement; values that do not converge are not quoted.",
    "• Only the verification problem is a thesis result. No experimental data are used.",
], size=20)

n = 3
s = frame("1. Verification: the thesis problem in 3D", n); n += 1
h = movie(s, "anim_wp3_seed_3d", Inches(0.7), Inches(1.0), Inches(11.9))
text(s, Inches(0.7), Inches(1.05) + h, Inches(11.9), Inches(1.6), [
    "β = 0.02 mm²/T*, k_α = 1e−3 /T*, E = 10 Pa, ν = 0.49, T* = 1.1; one octant cut away. Seed von Mises at T* = 1.1: "
    "6.19e−4 Pa (thesis, 16³: 6.18e−4 Pa). The stress forms a ring at the edge of the seed."], size=16)

s = frame("2. Stress next to a fixed surface: what converges", n); n += 1
s.shapes.add_picture(str(A / "fig_wall_convergence.png"), Inches(1.6), Inches(1.0), width=Inches(10.0))
text(s, Inches(0.7), Inches(5.7), Inches(11.9), Inches(1.6), [
    "Seed on a clamped face of the cube, meshes 16³ to 48³. The mean stress in the biofilm next to the face converges "
    "(less than 1 % change from 24³ on) and is 10 to 12 times the seed stress of the verification problem. The pointwise "
    "maximum at the corner grows as about h^−0.4, so I quote means and percentiles only."], size=16)

s = frame("3. A biofilm on a real molar", n); n += 1
h = movie(s, "anim_real_tooth_resolved_base_3d", Inches(0.7), Inches(1.0), Inches(11.9))
text(s, Inches(0.7), Inches(1.05) + h, Inches(11.9), Inches(1.8), [
    "Lower first molar at real size (crown about 10 × 12 mm) in a 16 mm box, root clamped; a 0.5 mm film on one flank at "
    "the cervical third. β scaled to 0.08 mm²/T* (an assumption). Growth converged to 2 %; stress from the film model."],
    size=16)

s = frame("4. Resolving the stress in a 0.5 mm film", n); n += 1
s.shapes.add_picture(str(A / "fig_real_tooth_submodel.png"), Inches(3.4), Inches(0.95), height=Inches(4.6))
text(s, Inches(0.7), Inches(5.6), Inches(11.9), Inches(1.8), [
    "On meshes of the whole box the film is one or two elements thick and the stress does not converge. I solve the film "
    "alone with the tooth rigid (0.1 % effect), down to h = 0.028 mm. With six elements across the film the means converge: "
    "film 3.3 mPa (±1 %), cervical band 3.5, the 0.1 mm next to the tooth 7.8, 95th percentile 10.5 mPa."], size=15)

s = frame("5. Nutrient and the growth limit on the molar", n); n += 1
s.shapes.add_picture(str(A / "fig_real_tooth_resolved.png"), Inches(1.9), Inches(0.95), width=Inches(9.5))
text(s, Inches(0.7), Inches(5.45), Inches(11.9), Inches(1.9), [
    "Nutrient from the saliva, consumption gφc with g/d = 10 /mm² (chosen; 10 times the Table 2 ratio): c falls to 0.07 at "
    "the tooth, growth and stress drop to about one seventh. At the Table 2 ratio they drop to 0.68. The growth limit alone "
    "(p_h half the film pressure without it, a choice) gives 0.71. Stress: end states re-solved with the film model."],
    size=15)

s = frame("5. Nutrient and growth limit: videos", n); n += 1
h = movie(s, "anim_real_tooth_resolved_nutrient_gate_3d", Inches(0.7), Inches(1.0), Inches(11.9))
text(s, Inches(0.7), Inches(1.05) + h, Inches(11.9), Inches(1.6), [
    "Nutrient and growth limit together; same colour scale for the stress as the base video (0 to 14.2 mPa)."], size=16)

s = frame("6. Four species on a tooth-like collar (synthetic data)", n); n += 1
h = movie(s, "anim_cervical_f4_3d", Inches(0.7), Inches(1.0), Inches(11.9))
text(s, Inches(0.7), Inches(1.05) + h, Inches(11.9), Inches(2), [
    "A 0.25 mm layer on a crown-shaped collar, inner face bonded, nutrient from the outside. Point model: the four-species "
    "case of Fritsch et al. (2025) with their true parameters. One quarter cut away; finite strain."], size=16)

s = frame("6. Section and composition", n); n += 1
s.shapes.add_picture(str(A / "fig_cervical_f4.png"), Inches(1.4), Inches(1.0), width=Inches(10.5))
text(s, Inches(0.7), Inches(5.75), Inches(11.9), Inches(1.6), [
    "(a) Species 3 takes over the seed (share 0.47), species 4 almost disappears (0.02). (b), (c) Section at 45°: without "
    "the limit the von Mises stress reaches 0.68 E at the bonded base; with p_h/E = 0.05 the largest value is 0.12 E."],
    size=16)

s = frame("Summary", n); n += 1
text(s, Inches(0.6), Inches(1.3), Inches(12.1), Inches(5.6), [
    "• The 3D view reproduces the thesis value of the seed stress (0.2 % difference).",
    "• Next to a fixed surface the biofilm stress is 10 to 12 times the free-seed value (converged band mean).",
    "• On a real molar a 0.5 mm film needs about six elements across it; then the film mean is 3.3 mPa (E = 10 Pa), "
    "highest in the 0.1 mm next to the tooth.",
    "• A nutrient limited by consumption lowers growth and stress more than the pressure limit.",
    "• Tooth surface: Diaz, Ferro Sánchez, Campo, Mendeley Data (2024), doi:10.17632/xjsx7nfhj8.1, CC BY 4.0.",
    "• Next at Keio: the calibrated parameters of Chapter 3 on this molar, and an implant."], size=20)

# speaker notes (EN) from speaker_notes_3d_demo.md, in slide order
import re
secs = re.split(r"^## \d+\. ", (ROOT / "speaker_notes_3d_demo.md").read_text(), flags=re.M)[1:]
notes = []
for sec in secs:
    m = re.search(r"EN: (.*?)\nJA:", sec, flags=re.S)
    notes.append(" ".join(m.group(1).split()) if m else "")
notes.insert(7, notes[6])                     # the second slide of case 5 (videos) repeats the note of the first
for sl, nt in zip(prs.slides, notes):
    sl.notes_slide.notes_text_frame.text = nt

out = ROOT / "slides_3d_demo.pptx"
prs.save(out)
print("wrote", out, "slides", len(prs.slides), f"{out.stat().st_size / 1e6:.1f} MB")
